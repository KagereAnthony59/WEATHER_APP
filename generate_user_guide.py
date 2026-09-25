import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_user_guide():
    doc = Document()

    # Set 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Color Palette
    COLOR_PRIMARY = RGBColor(26, 54, 93)      # Deep Navy #1A365D
    COLOR_SECONDARY = RGBColor(14, 116, 144)  # Ocean Blue #0E7490
    COLOR_ACCENT = RGBColor(217, 119, 6)      # Amber #D97706
    COLOR_DARK = RGBColor(30, 41, 59)         # Slate 800 #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500 #64748B
    COLOR_BG_LIGHT = "F8FAFC"
    COLOR_PRIMARY_HEX = "1A365D"
    COLOR_SECONDARY_HEX = "0E7490"

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Helper Functions
    def set_cell_shading(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.italic = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.bold = True
            run_b.font.color.rgb = COLOR_DARK
        p.add_run(text)
        return p

    def add_tip_box(text, title="PRO TIP"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_shading(cell, COLOR_BG_LIGHT)
        set_cell_margins(cell, 100, 100, 140, 140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run_t = p.add_run(f"💡 {title}: ")
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_SECONDARY
        p.add_run(text)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # DOCUMENT COVER & OVERVIEW
    # -------------------------------------------------------------
    add_title("K & A Weather App — User Guide")
    add_subtitle("Comprehensive Walkthrough, Biometeorology, Commute Safety, Features & Updates")

    add_tip_box(
        "Welcome to K & A Weather! This illustrated guide walks you through every feature of the application—from real-time weather forecasts and biometeorological health indices to driving safety analyses, 48-hour time travel, ambient soundscapes, and automatic Over-The-Air updates.",
        "WELCOME"
    )

    # -------------------------------------------------------------
    # 1. GETTING STARTED & APP NAVIGATION
    # -------------------------------------------------------------
    add_heading_1("1. Getting Started & Dashboard Overview")
    p = doc.add_paragraph()
    p.add_run(
        "When you launch K & A Weather, the app instantly hydrates your last-viewed location in under 50ms before seamlessly refreshing with the latest global satellite and meteorological data. "
        "The interface adapts dynamically with glassmorphic styling, condition-matched backgrounds, and smooth physics-driven particle animations."
    )

    add_heading_2("Top Floating Control Bar")
    add_bullet("Type any city, town, university campus, or region globally. Real-time autocomplete suggestions appear instantly.", "🔍 Search Bar: ")
    add_bullet("Instantly query your device's GPS hardware to fetch local weather for your exact location.", "📍 Current Location Button: ")
    add_bullet("Generate and share a beautiful weather graphic card to social media, messaging apps, or email.", "📤 Share Card Button: ")
    add_bullet("Open app settings to toggle between Celsius (°C) and Fahrenheit (°F) and review app version details.", "⚙️ Settings Button: ")
    add_bullet("Open the side-by-side comparison matrix for your saved cities.", "📊 Compare Button: ")

    # -------------------------------------------------------------
    # 2. INTELLIGENCE & ADVISORY SUITE
    # -------------------------------------------------------------
    add_heading_1("2. Intelligence, Wisdom & Daily Guidance")

    add_heading_2("AI Weather Narrative & Morning Briefings")
    p = doc.add_paragraph()
    p.add_run(
        "Located right below the top navigation, the Weather Narrative card delivers a natural-language contextual summary of the day. "
        "It automatically tailors messages for morning, afternoon, evening, and night, comparing today's temperature and humidity directly to yesterday so you know exactly what changed."
    )

    add_heading_2("Weather Wisdom Advisory Engine")
    p = doc.add_paragraph()
    p.add_run("A dedicated smart recommendation engine that analyzes real-feel temperature, wind chill, and UV to provide instant advice:")
    add_bullet("Actionable attire guidance (breathable cottons, light jackets, thermal layers, raincoats, or umbrellas).", "🧥 Outfit Recommendations: ")
    add_bullet("Hydration alerts, sunscreen reminders, and sudden pressure fluctuation warnings.", "🩺 Health & Wellness: ")
    add_bullet("Advance notices on upcoming rain windows, road spray, or low visibility for commuters.", "🚗 Commute & Travel: ")

    add_heading_2("Driving & Commute Safety Index")
    p = doc.add_paragraph()
    p.add_run(
        "Whether you're driving to work, riding a motorcycle, or commuting, the Driving Commute Safety card computes real-time roadway risk:\n"
        "• Road Risk Rating: Safe (Green), Moderate Caution (Amber), or Severe Caution (Red).\n"
        "• Hazard Breakdown: Calculates precipitation rate, road slickness factor, crosswind gust hazards, and atmospheric visibility.\n"
        "• Driver Advice: Recommends headlight usage, increased following distance, and hydroplaning precautions."
    )

    # -------------------------------------------------------------
    # 3. CELESTIAL ARC & BIOMETEOROLOGICAL SUITE
    # -------------------------------------------------------------
    add_heading_1("3. Celestial Arc & Health Intelligence")

    add_heading_2("Celestial Arc — Sun & Moon Tracker")
    p = doc.add_paragraph()
    p.add_run(
        "The Celestial Arc visualizes the real-time parabolic trajectory of the sun or moon across the horizon. "
        "It highlights critical solar milestones including Sunrise, Morning Golden Hour (ideal for photography), Solar Noon, Peak UV Hour, Sunset, Dusk, and Night Moon Phases."
    )

    add_heading_2("Comprehensive Biometeorology & Health Hub")
    p = doc.add_paragraph()
    p.add_run("Designed for wellness and environmental sensitivity, this suite tracks invisible atmospheric factors that impact health:")
    add_bullet("Forecasts botanical allergens (Grass, Birch, Weed pollen) to help allergy and asthma sufferers plan ahead.", "🌸 Pollen & Allergy Index: ")
    add_bullet("Monitors rapid drops in barometric pressure combined with high humidity that often trigger joint aches and arthritis flares.", "🦴 Joint Pain & Arthritis Sensitivity: ")
    add_bullet("Identifies sudden atmospheric pressure swings known to provoke vascular headaches and migraines.", "⚡ Migraine Risk Index: ")
    add_bullet("Displays current solar UV intensity along with safe unprotected sun exposure time estimates.", "☀️ UV Index & Sun Protection: ")
    add_bullet("Comprehensive European Air Quality Index with fine particulate (PM2.5), coarse dust (PM10), Ozone, NO2, and SO2 breakdowns.", "🍃 Air Quality (AQI): ")

    # -------------------------------------------------------------
    # 4. LIFESTYLE ADVISORIES & 48-HOUR TIME TRAVEL
    # -------------------------------------------------------------
    add_heading_1("4. Lifestyle Indices & Time Travel Scrubbing")

    add_heading_2("Lifestyle & Activity Suitability Index")
    p = doc.add_paragraph()
    p.add_run("Evaluates weather suitability scores (0–100%) for 6 popular everyday activities:")
    add_bullet("Scores based on temperature comfort, humidity, and air quality index.", "🏃 Running & Jogging: ")
    add_bullet("Evaluates wind speed, headwind gusts, and road wetness.", "🚴 Cycling & Biking: ")
    add_bullet("Calculates optimal night viewing based on cloud cover percentage, visibility, and moon phase.", "✨ Stargazing & Astronomy: ")
    add_bullet("Assesses patio and outdoor seating comfort based on breeze and rain probability.", "🍽️ Outdoor Dining: ")
    add_bullet("Estimates drying speed using solar radiation, wind speed, and humidity evaporation rates.", "🧺 Laundry Drying: ")
    add_bullet("Evaluates 48-hour precipitation probability to prevent rain right after washing your car.", "🚗 Car Washing: ")

    add_heading_2("48-Hour Interactive Time Travel Slider")
    p = doc.add_paragraph()
    p.add_run(
        "Drag the horizontal time scrubber forward up to 48 hours to preview upcoming weather. "
        "As you slide, the dashboard temperature, condition icons, wind indicators, and background lighting dynamically transition to simulate future conditions in real time."
    )

    # -------------------------------------------------------------
    # 5. RADAR MAPS, SOUNDSCAPES & SOCIAL SHARING
    # -------------------------------------------------------------
    add_heading_1("5. Interactive Radar, Soundscapes & Sharing")

    add_heading_2("Interactive Multi-Layer Weather Radar")
    p = doc.add_paragraph()
    p.add_run(
        "Tap the radar button to open a high-performance interactive Leaflet map. "
        "Switch seamlessly between Precipitation Radar, Temperature Heatmap, Wind Streamlines, and Cloud Satellites. Tap anywhere on the globe to inspect local conditions."
    )

    add_heading_2("Ambient Weather Soundscape Player")
    p = doc.add_paragraph()
    p.add_run(
        "Need to relax, focus, or sleep? The Soundscape Player streams high-fidelity ambient audio directly within the app. "
        "Choose from Rain on Roof, Gentle Wind, Mountain Breeze, Forest Birds, or Ocean Waves. Audio loops continuously in the background."
    )

    add_heading_2("Social Weather Share Card Generator")
    p = doc.add_paragraph()
    p.add_run(
        "Tap the share icon in the top bar to generate a high-resolution weather snapshot card featuring the city skyline backdrop, current metrics, and date stamp. "
        "Share directly to WhatsApp, Instagram Stories, Twitter/X, or save to your photo gallery."
    )

    # -------------------------------------------------------------
    # 6. AUTOMATIC UPDATES & OFFLINE CAPABILITY
    # -------------------------------------------------------------
    add_heading_1("6. Offline Resilience & Automatic OTA Updates")

    add_heading_2("Offline-First Functionality")
    p = doc.add_paragraph()
    p.add_run(
        "K & A Weather never leaves you with a blank screen. "
        "Your last forecast, saved cities, and preferences are securely cached locally, ensuring full access even on planes, in subways, or in areas with poor cellular reception."
    )

    add_heading_2("Seamless Over-The-Air (OTA) Updates")
    p = doc.add_paragraph()
    p.add_run(
        "You never have to manually download APK updates for new features or bug fixes. "
        "Whenever the engineering team publishes an update, the app automatically downloads and applies the latest version in the background upon launch."
    )

    add_tip_box(
        "Thank you for using K & A Weather! For questions, suggestions, or feedback, visit our GitHub repository or reach out to the development team.",
        "ENJOY"
    )

    output_path = r"d:\PROJECTS\WEATHER_APP\K_and_A_Weather_User_Guide.docx"
    doc.save(output_path)
    print(f"User Guide successfully created at: {output_path}")

if __name__ == "__main__":
    create_user_guide()
